import asyncio
import os
import sys
import shutil
import docx
import openpyxl
import pptx
from sqlalchemy import select

# Add apps/backend-ai to PYTHONPATH
sys.path.append(os.path.join(os.path.dirname(os.path.dirname(__file__)), "apps", "backend-ai"))

import app.models as models
from app.database import engine, AsyncSessionLocal
from app.services.document_ingestion import document_ingestion_service, STORAGE_DIR
from app.routers.document_agent import upload_document
from fastapi import UploadFile, BackgroundTasks

# Color constants
GREEN = '\x1b[32m'
RED = '\x1b[31m'
RESET = '\x1b[0m'
YELLOW = '\x1b[33m'

failed_tests = 0

def assert_test(condition, message):
    global failed_tests
    if condition:
        print(f"{GREEN}PASS:{RESET} {message}")
    else:
        print(f"{RED}FAIL:{RESET} {message}")
        failed_tests += 1

def generate_pdf(filepath, text):
    with open(filepath, "wb") as f:
        f.write(b'%PDF-1.4\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>\nendobj\n4 0 obj\n<< /Length ' + str(len(text) + 44).encode() + b' >>\nstream\nBT\n/F1 12 Tf\n72 712 Td\n(' + text.encode() + b') Tj\nET\nendstream\nendobj\nxref\n0 5\n0000000000 65535 f\n0000000009 00000 n\n0000000058 00000 n\n0000000115 00000 n\n0000000282 00000 n\ntrailer\n<< /Size 5 /Root 1 0 R >>\nstartxref\n377\n%%EOF')

def generate_docx(filepath, text):
    doc = docx.Document()
    doc.add_paragraph(text)
    doc.save(filepath)

def generate_xlsx(filepath, text):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = text
    wb.save(filepath)

def generate_pptx(filepath, text):
    prs = pptx.Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    left = top = width = height = pptx.util.Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = text
    prs.save(filepath)

async def test_file_conversions():
    print(f"\n{YELLOW}[TEST AREA 1] Real Binary Documents Conversion via MarkItDown{RESET}")
    
    # 1. Clean local storage directory
    if os.path.exists(STORAGE_DIR):
        shutil.rmtree(STORAGE_DIR)
    os.makedirs(STORAGE_DIR, exist_ok=True)
    
    temp_dir = os.path.join(os.path.dirname(__file__), "temp_binaries")
    os.makedirs(temp_dir, exist_ok=True)
    
    pdf_path = os.path.join(temp_dir, "test.pdf")
    docx_path = os.path.join(temp_dir, "test.docx")
    xlsx_path = os.path.join(temp_dir, "test.xlsx")
    pptx_path = os.path.join(temp_dir, "test.pptx")
    
    try:
        # Generate files
        generate_pdf(pdf_path, "Hello from dynamically generated PDF!")
        generate_docx(docx_path, "Hello from dynamically generated DOCX!")
        generate_xlsx(xlsx_path, "Hello from dynamically generated XLSX!")
        generate_pptx(pptx_path, "Hello from dynamically generated PPTX!")
        
        # Test PDF Ingestion
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()
        pdf_md = document_ingestion_service.convert_to_markdown(pdf_bytes, "test.pdf")
        assert_test("Hello from dynamically" in pdf_md, f"Successfully parsed PDF. Converted output snippet: '{pdf_md.strip()}'")
        
        # Test DOCX Ingestion
        with open(docx_path, "rb") as f:
            docx_bytes = f.read()
        docx_md = document_ingestion_service.convert_to_markdown(docx_bytes, "test.docx")
        assert_test("Hello from dynamically" in docx_md, f"Successfully parsed DOCX. Converted output snippet: '{docx_md.strip()}'")
        
        # Test XLSX Ingestion
        with open(xlsx_path, "rb") as f:
            xlsx_bytes = f.read()
        xlsx_md = document_ingestion_service.convert_to_markdown(xlsx_bytes, "test.xlsx")
        assert_test("Hello from dynamically" in xlsx_md, f"Successfully parsed XLSX. Converted output snippet: '{xlsx_md.strip()}'")
        
        # Test PPTX Ingestion
        with open(pptx_path, "rb") as f:
            pptx_bytes = f.read()
        pptx_md = document_ingestion_service.convert_to_markdown(pptx_bytes, "test.pptx")
        assert_test("Hello from dynamically" in pptx_md, f"Successfully parsed PPTX. Converted output snippet: '{pptx_md.strip()}'")
        
    finally:
        # Cleanup
        shutil.rmtree(temp_dir, ignore_errors=True)

class MockUploadFile(UploadFile):
    def __init__(self, filename: str, content: bytes):
        super().__init__(filename=filename, file=None)
        self._content = content
        self._cursor = 0

    async def read(self, size: int = -1) -> bytes:
        return self._content

    async def seek(self, offset: int) -> None:
        self._cursor = offset

async def test_concurrent_uploads():
    print(f"\n{YELLOW}[TEST AREA 2] 20 Concurrent Duplicate Uploads Stress-Test{RESET}")
    
    import uuid
    # Generate valid UTF-8 content with a unique UUID to prevent UnicodeDecodeError
    test_content = f"This is unique concurrent test file contents: {uuid.uuid4()}".encode("utf-8")
    filename = "concurrent_test_doc.txt"
    
    # Ensure RAG schema is updated
    async with engine.begin() as conn:
        await conn.run_sync(models.Base.metadata.create_all)
        
    # Helper async task for calling endpoint route with its own DB session
    async def task_call():
        async with AsyncSessionLocal() as db:
            bg_tasks = BackgroundTasks()
            upload_file = MockUploadFile(filename=filename, content=test_content)
            res = await upload_document(
                file=upload_file,
                commit_message="Concurrent upload test",
                background_tasks=bg_tasks,
                db=db
            )
            return res

    # Run 20 uploads simultaneously
    tasks = [task_call() for _ in range(20)]
    results = await asyncio.gather(*tasks, return_exceptions=True)
    
    # Assertions
    failures = [r for r in results if isinstance(r, BaseException)]
    for f in failures:
        import traceback
        print(f"{RED}Transaction Error:{RESET} {type(f).__name__}: {f}")
        if hasattr(f, "__traceback__"):
            traceback.print_exception(type(f), f, f.__traceback__)
        
    assert_test(len(failures) == 0, f"All 20 concurrent transactions completed successfully. Failures count: {len(failures)}")
    
    # Verify that all returned identical Document IDs (reused identical DB record)
    doc_ids = set()
    for res in results:
        if not isinstance(res, BaseException):
            doc_ids.add(res.id)
            
    assert_test(len(doc_ids) == 1, f"All concurrent calls correctly returned the same single Document ID: {list(doc_ids)}")
    
    # Verify DB records count
    async with AsyncSessionLocal() as db:
        docs = (await db.execute(select(models.Document))).scalars().all()
        versions = (await db.execute(select(models.Version))).scalars().all()
        
        assert_test(len(docs) == 1, f"Database contains exactly 1 Document record (No duplicates created). Actual count: {len(docs)}")
        assert_test(len(versions) == 1, f"Database contains exactly 1 Version record. Actual count: {len(versions)}")
        
        # Cleanup
        if len(docs) > 0:
            doc_id = docs[0].id
            await db.execute(models.Document.__table__.delete().where(models.Document.id == doc_id))
            await db.commit()
            
            # Check local file cleanup
            storage_file = os.path.join(STORAGE_DIR, f"{doc_id}.md")
            if os.path.exists(storage_file):
                os.remove(storage_file)
                print(f"Cleaned local storage file: {storage_file}")

async def main():
    print("\n==================================================================")
    print("STARTING ADVANCED PRODUCTION INGESTION & CONCURRENCY TESTS")
    print("==================================================================")
    
    try:
        await test_file_conversions()
        await test_concurrent_uploads()
    except Exception as err:
        print(f"{RED}Test runner crashed: {err}{RESET}")
        global failed_tests
        failed_tests += 1

    print("\n==================================================================")
    print("PRODUCTION TESTING REPORT")
    print("==================================================================")
    if failed_tests == 0:
        print(f"{GREEN}ALL ADVANCED INGESTION AND CONCURRENCY TESTS PASSED SUCCESSFULLY!{RESET}\n")
        sys.exit(0)
    else:
        print(f"{RED}TEST SUITE EXITED WITH {failed_tests} FAILURES{RESET}\n")
        sys.exit(1)

if __name__ == "__main__":
    asyncio.run(main())
